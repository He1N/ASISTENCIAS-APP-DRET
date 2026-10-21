from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
from unittest.mock import patch
from django.contrib.auth.models import Group, Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from .importacion import (ErrorImportacion, anular_lote, confirmar, crear_borrador, octubre_historico,
    preparar_plan, revisar, seleccionar_hoja)
from .models import (HistorialCambio, IncidenciaImportacion, Justificacion, LoteImportacion,
    PeriodoAsistencia, TipoDocumento, TipoOperacion, Trabajador)
from .xlsx_parser import ErrorXLSX, leer_archivo


def archivo(rows=None, nombre='prueba.xlsx', hoja='REGISTROS', reorder=False, encabezado=2):
    wb=Workbook(); ws=wb.active; ws.title=hoja
    headers=['Items','Apellidos y Nombres','Tipo Documento','Tipo Operación','Fecha Movimiento',
             'Salida Hora','Salida Minuto','Retorno Hora','Retorno Minuto','Observaciones','Fecha Final','N° CUD']
    order=list(range(12)) if not reorder else [2,4,1,3,0,11,5,6,7,8,10,9]
    for idx,col in enumerate(order,1): ws.cell(encabezado,idx,headers[col])
    for num,row in enumerate(rows or [fila()],encabezado+3):
        for idx,col in enumerate(order,1): ws.cell(num,idx,row[col])
    wb.create_sheet('OJO'); stream=BytesIO();wb.save(stream)
    return SimpleUploadedFile(nombre,stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def fila(nombre='PERSONA FICTICIA | 00123456', cud='00-ABC / 002', item=1, fecha=date(2026,10,1), final=date(2026,10,1), operacion='Operación ficticia | 001'):
    return [item,nombre,'Papeleta de Permiso | 01',operacion,fecha,10,0,11,30,'Observación ficticia',final,cud]


@override_settings(DEBUG=False, SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False, CSRF_COOKIE_SECURE=False)
class ExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('preparar_roles',stdout=StringIO())
        cls.users={}
        for rol in ('Administrador','Operador','Consulta'):
            u=User.objects.create_user(rol.lower());u.groups.add(Group.objects.get(name=rol));cls.users[rol]=u
        cls.op=TipoOperacion.objects.create(nombre='Operación ficticia',codigo_externo='001',requiere_fecha_final=True)
        cls.oct=PeriodoAsistencia.objects.get(titulo='Octubre 2026')
        cls.nov=PeriodoAsistencia.objects.get(titulo='Noviembre 2026')

    def login(self,rol='Operador'): self.client.force_login(self.users[rol])

    def borrador(self,rows=None,periodo=None,**kwargs):
        file=archivo(rows,**kwargs)
        if rows is None and not kwargs:
            if hasattr(self, '_archivo_default'):
                file=SimpleUploadedFile('prueba.xlsx', self._archivo_default)
            else:
                self._archivo_default=file.read();file.seek(0)
        parsed=leer_archivo(file)
        lote=crear_borrador(parsed,'prueba.xlsx',periodo or self.oct,self.users['Operador'])
        return seleccionar_hoja(lote.pk,kwargs.get('hoja','REGISTROS'))

    def revision(self,lote,choices=None,permitir=False):
        return revisar(lote.pk,choices or {},permitir)

    def confirmar(self,lote,choices=None,permitir=False):
        lote,plan=self.revision(lote,choices,permitir)
        return confirmar(lote.pk,self.users['Operador'],lote.datos_temporales['revision'])

    def test_parseo_valido_mapeo_reordenado_y_ceros(self):
        parsed=leer_archivo(archivo(reorder=True))
        self.assertNotIn('OJO',parsed['hojas'])
        self.assertEqual(parsed['hojas']['REGISTROS']['mapeo']['trabajador'],'C')
        lote=crear_borrador(parsed,'prueba.xlsx',self.oct,self.users['Operador'])
        lote=seleccionar_hoja(lote.pk,'REGISTROS')
        self.confirmar(lote)
        obj=Justificacion.objects.get()
        self.assertEqual(obj.trabajador.dni,'00123456')
        self.assertEqual(obj.tipo_documento.codigo_externo,'01')
        self.assertEqual(obj.fecha_movimiento,date(2026,10,1))
        self.assertEqual(obj.hora_salida,time(10,0));self.assertEqual(obj.hora_retorno,time(11,30))
        self.assertEqual(obj.fila_origen,5);self.assertEqual(obj.item_origen,'1')

    def test_octubre_descarta_todos_cud_noviembre_conserva(self):
        lote=self.borrador([fila(cud='SOSPECHOSO')]);self.assertNotIn('SOSPECHOSO',str(lote.datos_temporales))
        self.confirmar(lote);obj=Justificacion.objects.get();self.assertIsNone(obj.cud)
        self.assertIn('cud_omitido_por_desalineacion_octubre',lote.reglas_aplicadas)
        rows=[fila(fecha=date(2026,10,15),final=date(2026,10,15),cud='00-CUD-XYZ / 003')]
        lote=self.borrador(rows,periodo=self.nov);self.confirmar(lote)
        self.assertEqual(Justificacion.objects.latest('pk').cud,'00-CUD-XYZ / 003')
        self.assertIsNone(Justificacion.objects.first().fecha_recepcion)

    def test_excepcion_por_corte_no_por_mes_movimiento_ni_otros_anios(self):
        p=PeriodoAsistencia.objects.create(titulo='Octubre 2027',fecha_inicio=date(2027,9,7),fecha_fin=date(2027,10,9))
        self.assertFalse(octubre_historico(p));self.assertTrue(octubre_historico(self.oct))
        lote=self.borrador([fila(fecha=date(2027,10,1),final=date(2027,10,1),cud='CONSERVAR')],periodo=p)
        self.confirmar(lote);self.assertEqual(Justificacion.objects.get().cud,'CONSERVAR')

    def test_no_modifica_cud_existente_y_duplicados_requieren_revision(self):
        worker=Trabajador.objects.create(apellidos_nombres='PERSONA FICTICIA ',dni='00123456')
        previo=Justificacion.objects.create(periodo=self.oct,trabajador=worker,tipo_documento=TipoDocumento.objects.get(codigo_externo='01'),tipo_operacion=self.op,fecha_movimiento=date(2026,10,1),created_by=self.users['Operador'],updated_by=self.users['Operador'],cud='EXISTENTE')
        lote=self.borrador();plan=preparar_plan(lote)
        self.assertEqual(plan['totales']['duplicadas'],1);self.assertEqual(plan['totales']['listas'],0)
        with self.assertRaises(ErrorImportacion): self.confirmar(lote)
        choices={'5':{'accion':'confirmar_duplicado'}}
        self.confirmar(lote,choices)
        previo.refresh_from_db();self.assertEqual(previo.cud,'EXISTENTE')
        self.assertEqual(Justificacion.objects.count(),2)

    def test_reimportacion_y_confirmacion_doble_no_duplican(self):
        lote=self.confirmar(self.borrador())
        repetir=self.borrador();self.assertEqual(repetir.pk,lote.pk)
        confirmar(lote.pk,self.users['Operador'],'cualquier-version')
        self.assertEqual(Justificacion.objects.count(),1)
        self.assertEqual(Trabajador.objects.count(),1)

    def test_duplicado_dentro_archivo_no_borra_ni_asume_unicidad(self):
        lote=self.borrador([fila(item=1),fila(item=2)])
        plan=preparar_plan(lote);self.assertEqual(plan['totales']['duplicadas'],1)
        self.confirmar(lote)
        self.assertEqual(Justificacion.objects.count(),1)
        self.assertEqual(LoteImportacion.objects.get(pk=lote.pk).omitidos,1)

    def test_homonimos_sin_dni_se_resuelven_por_id(self):
        a=Trabajador.objects.create(apellidos_nombres='PERSONA FICTICIA')
        b=Trabajador.objects.create(apellidos_nombres='PERSONA FICTICIA')
        lote=self.borrador([fila(nombre='PERSONA FICTICIA')]);plan=preparar_plan(lote)
        self.assertEqual(plan['totales']['listas'],0)
        self.assertNotIn('trabajador',plan['filas'][0]['elegidos'])
        self.confirmar(lote,{'5':{'trabajador':str(b.pk)}})
        self.assertEqual(Justificacion.objects.get().trabajador_id,b.pk)
        self.assertEqual(Trabajador.objects.count(),2)

    def test_nombre_solo_y_dni_invalido_no_se_inventa(self):
        lote=self.borrador([fila(nombre='SOLO NOMBRE FICTICIO')]);self.confirmar(lote)
        self.assertIsNone(Trabajador.objects.get().dni)
        self.assertEqual(Trabajador.objects.get().apellidos_nombres,'SOLO NOMBRE FICTICIO')

    def test_pendiente_y_faltantes_son_incidencias_no_valores_inventados(self):
        self.op.requiere_fecha_final=None;self.op.save()
        lote=self.borrador([fila(fecha=None,final=None)])
        plan=preparar_plan(lote);self.assertEqual(plan['totales']['incompletas'],1)
        self.assertEqual(plan['totales']['listas'],0)
        self.confirmar(lote,permitir=True)
        obj=Justificacion.objects.get();self.assertIsNone(obj.fecha_movimiento)
        self.assertEqual(obj.estado_calidad,'INCOMPLETO_LEGADO')
        self.assertIsNone(obj.es_registro_tardio)
        self.assertTrue(IncidenciaImportacion.objects.filter(lote=lote).exists())

    def test_fechas_contradictorias_y_horas_incompletas_bloquean_fila(self):
        row=fila(final=date(2026,9,1));row[6]=None
        lote=self.borrador([row]);plan=preparar_plan(lote,permitir_incompletos=True)
        self.assertEqual(plan['totales']['errores'],1)
        self.assertTrue(any('anterior' in i['mensaje'] for i in plan['filas'][0]['issues']))
        with self.assertRaises(ErrorImportacion): self.confirmar(lote,permitir=True)
        self.assertFalse(Justificacion.objects.exists())

    def test_periodo_cerrado_y_movimiento_fuera_rango(self):
        self.oct.estado='PROCESADO';self.oct.fecha_cierre=timezone.now();self.oct.save()
        lote=self.borrador([fila(fecha=date(2026,10,15),final=date(2026,10,15))])
        self.assertEqual(preparar_plan(lote)['totales']['listas'],0)
        self.confirmar(lote,{'5':{'motivo':'Corte administrativo ficticio autorizado'}})
        self.oct.refresh_from_db();self.assertEqual(self.oct.estado,'PROCESADO')
        obj=Justificacion.objects.get();self.assertEqual(obj.fecha_movimiento,date(2026,10,15))
        self.assertIsNone(obj.es_registro_tardio)

    def test_error_transaccional_no_deja_datos_parciales(self):
        lote=self.borrador([fila(item=1),fila(nombre='OTRA PERSONA FICTICIA',item=2)])
        lote,_=self.revision(lote)
        original=Justificacion.save;contador=[0]
        def fallo(obj,*args,**kwargs):
            contador[0]+=1
            if contador[0]==2: raise RuntimeError('fallo sintético')
            return original(obj,*args,**kwargs)
        with patch.object(Justificacion,'save',fallo):
            with self.assertRaises(RuntimeError): confirmar(lote.pk,self.users['Operador'],lote.datos_temporales['revision'])
        self.assertEqual(Justificacion.objects.count(),0);self.assertEqual(Trabajador.objects.count(),0)
        self.assertEqual(HistorialCambio.objects.count(),0)
        lote.refresh_from_db();self.assertEqual(lote.estado,'PREVIA')

    def test_cambios_maestro_obligan_revisar_y_version_obsoleta(self):
        lote,_=self.revision(self.borrador());version=lote.datos_temporales['revision']
        self.op.nombre='Nombre corregido';self.op.save()
        with self.assertRaises(ErrorImportacion): confirmar(lote.pk,self.users['Operador'],version)
        self.assertFalse(Justificacion.objects.exists())
        lote,_=self.revision(lote,{'5':{'operacion':str(self.op.pk)}})
        with self.assertRaises(ErrorImportacion): confirmar(lote.pk,self.users['Operador'],version)

    def test_anular_lote_sin_borrar_y_no_sobre_ediciones_posteriores(self):
        lote=self.confirmar(self.borrador())
        anular_lote(lote.pk,self.users['Administrador'],'Prueba de anulación de lote')
        self.assertTrue(Justificacion.objects.get().anulada)
        self.assertEqual(Justificacion.objects.count(),1)
        self.assertEqual(LoteImportacion.objects.get(pk=lote.pk).estado,'ANULADO')
        lote=self.confirmar(self.borrador([fila(nombre='OTRA PERSONA FICTICIA')]))
        obj=lote.justificacion_set.get();obj.observaciones='Corrección posterior';obj.save()
        with self.assertRaises(ErrorImportacion): anular_lote(lote.pk,self.users['Administrador'],'No autorizado por cambios posteriores')
        obj.refresh_from_db();self.assertFalse(obj.anulada)

    def test_archivos_invalidos_y_formulas_no_se_ejecutan(self):
        for file in (SimpleUploadedFile('x.xlsx',b'no soy un zip'),SimpleUploadedFile('x.xlsm',b'abc')):
            with self.assertRaises(ErrorXLSX): leer_archivo(file)
        row=fila(cud='=1+1')
        lote=self.borrador([row],periodo=self.nov)
        plan=preparar_plan(lote)
        self.assertTrue(any('Fórmula' in i['mensaje'] for i in plan['filas'][0]['issues']))
        # Octubre descarta hasta CUD con fórmula, sin adjudicarlo ni evaluarlo.
        lote=self.borrador([row]);self.confirmar(lote)
        self.assertIsNone(Justificacion.objects.get().cud)

    def test_lote_cancelado_se_puede_cargar_de_nuevo_y_purgar(self):
        lote=self.borrador();lote.estado='CANCELADO';lote.datos_temporales={};lote.save()
        nuevo=self.borrador();self.assertEqual(nuevo.pk,lote.pk)
        self.assertEqual(nuevo.estado,'PREVIA')
        LoteImportacion.objects.filter(pk=lote.pk).update(fecha=timezone.now()-timedelta(days=2))
        call_command('purgar_previsualizaciones',stdout=StringIO())
        lote.refresh_from_db();self.assertEqual(lote.estado,'EXPIRADO');self.assertEqual(lote.datos_temporales,{})

    def test_permisos_idor_y_csrf_importacion(self):
        lote=self.borrador()
        self.login('Consulta')
        for url in (reverse('importar_excel'),reverse('importacion_revisar',args=[lote.pk]),reverse('importacion_resumen',args=[lote.pk])):
            self.assertEqual(self.client.get(url).status_code,403)
        self.login('Administrador')
        self.assertEqual(self.client.get(reverse('importar_excel')).status_code,200)
        u=User.objects.create_user('otro_operador');u.groups.add(Group.objects.get(name='Operador'));self.client.force_login(u)
        self.assertEqual(self.client.get(reverse('importacion_revisar',args=[lote.pk])).status_code,403)
        c=Client(enforce_csrf_checks=True);c.force_login(self.users['Operador'])
        self.assertEqual(c.post(reverse('importar_excel'),{}).status_code,403)

    def test_flujo_web_completo_y_purga_al_confirmar(self):
        self.login()
        response=self.client.post(reverse('importar_excel'),{'archivo':archivo(),'periodo':self.oct.pk,'encabezado':2,'primera_fila':5})
        self.assertEqual(response.status_code,302);lote=LoteImportacion.objects.get(estado='SELECCION')
        url=reverse('importacion_revisar',args=[lote.pk])
        self.assertContains(self.client.get(url),'Seleccionar hoja')
        self.client.post(url,{'hoja':'REGISTROS','conservar_cud':'on'})
        response=self.client.get(url);self.assertContains(response,'Previsualización de importación')
        plan=response.context['plan'];row=plan['filas'][0]
        post={f'{key}_5':row['elegidos'].get(key,'') for key in ('trabajador','documento','operacion')}
        post.update(accion_5='importar',motivo_5='',accion='revisar')
        response=self.client.post(url,post);self.assertContains(response,'Confirmar importación revisada')
        post.update(accion='confirmar',revision=response.context['revision'])
        self.assertEqual(self.client.post(url,post).status_code,302)
        lote.refresh_from_db();self.assertEqual(lote.estado,'CONFIRMADO');self.assertEqual(lote.datos_temporales,{})
        self.assertContains(self.client.get(reverse('importacion_resumen',args=[lote.pk])),'Insertadas: 1')

    def test_cambiar_decision_al_confirmar_solo_revisa(self):
        self.login();lote,plan=self.revision(self.borrador())
        post={f'{key}_5':plan['filas'][0]['elegidos'].get(key,'') for key in ('trabajador','documento','operacion')}
        post.update(accion_5='importar',motivo_5='',accion='confirmar',revision=lote.datos_temporales['revision'])
        response=self.client.post(reverse('importacion_revisar',args=[lote.pk]),post)
        self.assertContains(response,'Las decisiones cambiaron')
        self.assertEqual(Justificacion.objects.count(),0)

    def test_511_filas_sinteticas_migradas_sin_pii_real(self):
        rows=[fila(nombre=f'PERSONA FICTICIA {n:04}',item=n) for n in range(511)]
        lote=self.confirmar(self.borrador(rows))
        self.assertEqual(lote.insertados,511);self.assertEqual(Justificacion.objects.count(),511)
        self.assertEqual(Trabajador.objects.count(),511)

    def test_encabezado_original_dos_filas_y_grupos_hora_minuto(self):
        wb=Workbook();ws=wb.active;ws.title='ORIGINAL'
        base=['Items','Apellidos y Nombres','Tipo Documento','Tipo Operación','Fecha Movimiento']
        for n,label in enumerate(base,1):ws.cell(2,n,label)
        ws['F2']='Salida';ws.merge_cells('F2:G2');ws['H2']='Retorno';ws.merge_cells('H2:I2')
        for col,label in [('F','Hora'),('G','Minuto'),('H','Hora'),('I','Minuto')]:ws[f'{col}3']=label
        ws['J2']='Observaciones';ws['K2']='Fecha Final';ws['L2']='N° CUD'
        for n,value in enumerate(fila(),1):ws.cell(5,n,value)
        stream=BytesIO();wb.save(stream)
        parsed=leer_archivo(SimpleUploadedFile('original-sintetico.xlsx',stream.getvalue()))
        mapping=parsed['hojas']['ORIGINAL']['mapeo']
        self.assertEqual(mapping['salida_hora'],'F');self.assertEqual(mapping['salida_minuto'],'G')
        self.assertEqual(mapping['retorno_minuto'],'I')

    def test_epoca_1904_y_cud_numerico_con_formato_conservan_datos(self):
        from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
        wb=Workbook();wb.epoch=CALENDAR_MAC_1904;ws=wb.active;ws.title='REGISTROS'
        headers=['Items','Apellidos y Nombres','Tipo Documento','Tipo Operación','Fecha Movimiento','Salida Hora','Salida Minuto','Retorno Hora','Retorno Minuto','Observaciones','Fecha Final','N° CUD']
        for n,label in enumerate(headers,1):ws.cell(2,n,label)
        row=fila(fecha=to_excel(datetime(2026,10,15),wb.epoch),final=to_excel(datetime(2026,10,15),wb.epoch),cud=7)
        for n,value in enumerate(row,1):ws.cell(5,n,value)
        ws['L5'].number_format='00000';stream=BytesIO();wb.save(stream)
        parsed=leer_archivo(SimpleUploadedFile('epoca.xlsx',stream.getvalue()))
        lote=crear_borrador(parsed,'epoca.xlsx',self.nov,self.users['Operador']);lote=seleccionar_hoja(lote.pk,'REGISTROS')
        self.confirmar(lote)
        obj=Justificacion.objects.get();self.assertEqual(obj.cud,'00007');self.assertEqual(obj.fecha_movimiento,date(2026,10,15))

    def test_asignacion_manual_no_cambia_catalogo_y_deja_incidencia(self):
        lote=self.borrador([fila(operacion='Operación ficticia | 999')])
        self.assertEqual(preparar_plan(lote)['totales']['listas'],0)
        self.confirmar(lote,{'5':{'operacion':str(self.op.pk)}})
        self.op.refresh_from_db();self.assertEqual(self.op.codigo_externo,'001')
        self.assertTrue(IncidenciaImportacion.objects.filter(lote=lote,valor_original__contains='999').exists())

    def test_celda_nombre_vacia_se_informa_sin_error_interno(self):
        row=fila();row[1]=None
        lote=self.borrador([row]);plan=preparar_plan(lote)
        self.assertEqual(plan['totales']['errores'],1)
        self.assertEqual(plan['totales']['listas'],0)
