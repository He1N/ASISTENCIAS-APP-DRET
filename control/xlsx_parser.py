"""Lectura acotada de XLSX. No evalúa fórmulas ni conserva el archivo original."""
import hashlib
import math
import re
from datetime import date, datetime, time
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from .models import normalizar

MAX_ARCHIVO = 10 * 1024 * 1024
MAX_DESCOMPRIMIDO = 50 * 1024 * 1024
MAX_FILAS = 5000
CAMPOS = {
    'item': 'Ítem', 'trabajador': 'Apellidos y nombres', 'documento': 'Tipo de documento',
    'operacion': 'Tipo de operación', 'movimiento': 'Fecha de movimiento',
    'salida_hora': 'Salida hora', 'salida_minuto': 'Salida minuto',
    'retorno_hora': 'Retorno hora', 'retorno_minuto': 'Retorno minuto',
    'observaciones': 'Observaciones', 'final': 'Fecha final', 'cud': 'CUD',
}
ALIASES = {
    'item': ['items', 'item', 'nro item'], 'trabajador': ['apellidos y nombres', 'apellidos nombres', 'trabajador'],
    'documento': ['tipo documento', 'tipo de documento'], 'operacion': ['tipo operacion', 'tipo de operacion'],
    'movimiento': ['fecha movimiento', 'fecha de movimiento', 'fecha inicial'],
    'salida_hora': ['salida hora', 'hora salida', 'hora de salida'], 'salida_minuto': ['salida minuto', 'minuto salida'],
    'retorno_hora': ['retorno hora', 'hora retorno', 'hora de retorno'], 'retorno_minuto': ['retorno minuto', 'minuto retorno'],
    'observaciones': ['observaciones', 'observacion'], 'final': ['fecha final', 'fecha fin'],
    'cud': ['n° cud', 'nº cud', 'n cud', 'no cud', 'cud', 'numero cud'],
}


class ErrorXLSX(ValueError):
    pass


def texto(celda):
    """Preserva códigos numéricos con formato explícito 000... cuando sea posible."""
    if celda.value is None:
        return None
    value = celda.value
    if celda.data_type == 'f':
        return {'formula': True}
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)):
            raise ErrorXLSX('El archivo contiene un valor numérico no finito.')
        if float(value).is_integer() and re.fullmatch(r'0+', celda.number_format or ''):
            return str(int(value)).zfill(len(celda.number_format))
    return value


def normalizar_encabezado(value):
    return normalizar(str(value or '')).replace('.', '').replace('_', ' ')


def mapeo_hoja(ws, fila_encabezado, overrides):
    # Solo se inspeccionan encabezados hasta BL. Datos ajenos al mapa no se leen.
    rows = list(ws.iter_rows(min_row=fila_encabezado, max_row=fila_encabezado + 1, max_col=64))
    top = [normalizar_encabezado(c.value) for c in rows[0]]
    bottom = [normalizar_encabezado(c.value) for c in rows[1]]
    mapping, grupo = {}, ''
    for idx, (a, b) in enumerate(zip(top, bottom), 1):
        if a in ('salida', 'retorno'):
            grupo = a
        elif a:
            grupo = ''
        label = f'{grupo} {b}' if grupo and b else a
        for campo, aliases in ALIASES.items():
            if overrides.get(campo):
                continue
            if label in aliases or (a and b and f'{a} {b}' in aliases):
                if campo in mapping:
                    raise ErrorXLSX(f'Encabezado ambiguo para {CAMPOS[campo]}; indique su columna explícita.')
                mapping[campo] = idx
    for campo, col in overrides.items():
        if col:
            mapping[campo] = column_index_from_string(col.upper())
    if len(set(mapping.values())) != len(mapping):
        raise ErrorXLSX('No asigne una columna a dos campos distintos.')
    for campo in ('trabajador', 'documento', 'operacion'):
        if campo not in mapping:
            raise ErrorXLSX(f'No se reconoció {CAMPOS[campo]}. Indique su columna en el mapeo.')
    return mapping


def leer_archivo(archivo, encabezado=2, primera_fila=5, overrides=None):
    if not archivo.name.lower().endswith('.xlsx'):
        raise ErrorXLSX('Seleccione un archivo .xlsx, sin macros.')
    raw = archivo.read(MAX_ARCHIVO + 1)
    if len(raw) > MAX_ARCHIVO:
        raise ErrorXLSX('El límite del archivo es 10 MB.')
    try:
        with ZipFile(BytesIO(raw)) as z:
            info = z.infolist()
            if len(info) > 2000 or sum(i.file_size for i in info) > MAX_DESCOMPRIMIDO:
                raise ErrorXLSX('El contenido descomprimido excede los límites de seguridad.')
            if any('vbaproject' in i.filename.lower() for i in info):
                raise ErrorXLSX('No se admiten macros.')
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        hojas, total = {}, 0
        for ws in wb.worksheets:
            if len(ws.title) > 100:
                raise ErrorXLSX('El nombre de una hoja supera 100 caracteres.')
            if normalizar(ws.title) == 'ojo':
                continue
            if len(hojas) >= 12:
                raise ErrorXLSX('Se admiten como máximo 12 hojas de registros por archivo.')
            try:
                mapping = mapeo_hoja(ws, encabezado, overrides or {})
            except ErrorXLSX as exc:
                hojas[ws.title] = {'error': str(exc), 'filas': [], 'mapeo': {}}
                continue
            filas = []
            # No se confía en dimensiones declaradas de un documento externo.
            ws.reset_dimensions()
            for number, row in enumerate(ws.iter_rows(min_row=primera_fila, max_col=max(mapping.values())), primera_fila):
                if number - primera_fila >= 20000:
                    raise ErrorXLSX('Se exceden 20000 filas inspeccionadas; recorte el área de datos.')
                values = {campo: texto(row[col - 1]) for campo, col in mapping.items()}
                if all(v is None or v == '' for v in values.values()):
                    continue
                for key in ('movimiento', 'final'):
                    value = values.get(key)
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        try:
                            parsed_date = from_excel(value, epoch=wb.epoch)
                            if isinstance(parsed_date, datetime):
                                values[key] = parsed_date.date().isoformat()
                        except (ValueError, OverflowError):
                            pass
                total += 1
                if total > MAX_FILAS:
                    raise ErrorXLSX('El límite es 5000 filas no vacías por archivo.')
                filas.append({'fila': number, 'raw': values})
            hojas[ws.title] = {'filas': filas, 'mapeo': {k: get_column_letter(v) for k, v in mapping.items()}}
        wb.close()
    except ErrorXLSX:
        raise
    except (BadZipFile, ValueError, KeyError, OSError, TypeError, IndexError, ParseError, DefusedXmlException, InvalidFileException) as exc:
        raise ErrorXLSX('No se pudo leer el contenido XLSX o su estructura.') from exc
    if not hojas or not any(h['filas'] for h in hojas.values()):
        raise ErrorXLSX('No hay hojas de registros reconocidas. Revise filas de encabezado, inicio y mapeo.')
    return {'checksum': hashlib.sha256(raw).hexdigest(), 'hojas': hojas}


def fecha(value):
    if value is None or value == '':
        return None
    if isinstance(value, dict):
        raise ValueError('La fecha contiene una fórmula; no se evalúa.')
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        parsed = from_excel(value)
        if isinstance(parsed, datetime):
            return parsed.date().isoformat()
        raise ValueError('Número de fecha Excel inválido.')
    try:
        return datetime.fromisoformat(str(value).strip()).date().isoformat()
    except ValueError:
        pass
    for fmt in ('%d/%m/%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Fecha inválida. Revise el dato original.')


def hora(h, m):
    if h in (None, '') and m in (None, ''):
        return None
    if isinstance(h, dict) or isinstance(m, dict):
        raise ValueError('La hora contiene una fórmula; no se evalúa.')
    if isinstance(h, str) and ':' in h and m in (None, ''):
        try:
            return time.fromisoformat(h).isoformat()
        except ValueError:
            raise ValueError('Hora inválida.')
    if h in (None, '') or m in (None, ''):
        raise ValueError('Hora/minuto incompletos; no se inventa el componente faltante.')
    try:
        if float(h) != int(float(h)) or float(m) != int(float(m)):
            raise ValueError
        return time(int(float(h)), int(float(m))).isoformat()
    except (ValueError, TypeError, OverflowError):
        raise ValueError('Hora o minuto inválidos.')


def separar(value):
    if not isinstance(value, str):
        return str(value or ''), ''
    if '|' not in value:
        return value, ''
    nombre, codigo = value.rsplit('|', 1)
    return nombre, codigo.strip()


def transformar(fila, omitir_cud=False):
    raw, issues = fila['raw'], []
    def incidencia(campo, mensaje, severidad='ALERTA', original=None):
        issues.append({'campo': campo, 'mensaje': mensaje, 'severidad': severidad,
                       'original': str(original)[:250] if original is not None else ''})
    values = {}
    nombre_raw = raw.get('trabajador')
    nombre, dni = separar(nombre_raw)
    if dni and not re.fullmatch(r'[0-9]{8}', dni):
        incidencia('trabajador', 'Delimitador sin DNI plausible: se conserva el texto completo.', original=nombre_raw)
        nombre, dni = str(nombre_raw), ''
    values.update(nombre=nombre, dni=dni or None)
    for campo in ('documento', 'operacion'):
        n, c = separar(raw.get(campo))
        values[campo] = {'nombre': n, 'codigo': c}
    if not nombre.strip():
        incidencia('trabajador', 'Falta nombre de trabajador.', 'ERROR')
    for source, target in [('movimiento', 'fecha_movimiento'), ('final', 'fecha_final')]:
        try:
            values[target] = fecha(raw.get(source))
        except (ValueError, OverflowError) as exc:
            values[target] = None
            incidencia(source, str(exc), 'ERROR', raw.get(source))
    if not values['fecha_movimiento']:
        incidencia('movimiento', 'Movimiento ausente: registro INCOMPLETO_LEGADO.')
    if values['fecha_movimiento'] and values['fecha_final'] and values['fecha_final'] < values['fecha_movimiento']:
        incidencia('final', 'Fecha final anterior al inicio; omita/corrija la fila en el archivo.', 'ERROR', raw.get('final'))
    for tipo in ('salida', 'retorno'):
        try:
            values['hora_' + tipo] = hora(raw.get(tipo + '_hora'), raw.get(tipo + '_minuto'))
        except ValueError as exc:
            values['hora_' + tipo] = None
            incidencia(tipo, str(exc), 'ERROR', [raw.get(tipo + '_hora'), raw.get(tipo + '_minuto')])
    for key in ('observaciones', 'item', 'cud'):
        value = raw.get(key)
        if key == 'cud' and omitir_cud:
            values[key] = None  # Ni siquiera se interpreta el CUD sospechoso.
        elif isinstance(value, dict):
            values[key] = None
            incidencia(key, 'Fórmula no admitida. Corrija u omita la fila.', 'ERROR')
        else:
            values[key] = str(value) if value is not None else None
    if isinstance(nombre_raw, dict) or any(isinstance(raw.get(c), dict) for c in ('documento', 'operacion')):
        incidencia('identidad', 'No se admiten fórmulas como nombres o catálogos.', 'ERROR')
    values['issues'] = issues
    return values
