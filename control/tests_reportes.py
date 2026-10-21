from datetime import date, time
from io import BytesIO, StringIO
from django.contrib.auth.models import Group, Permission, User
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from .models import Justificacion, PeriodoAsistencia, TipoDocumento, TipoOperacion, Trabajador


@override_settings(DEBUG=False, SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False, CSRF_COOKIE_SECURE=False)
class ReportesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('preparar_roles',stdout=StringIO())
        cls.user=User.objects.create_user('operador');cls.user.groups.add(Group.objects.get(name='Operador'))
        cls.reader=User.objects.create_user('consulta');cls.reader.groups.add(Group.objects.get(name='Consulta'))
        cls.worker=Trabajador.objects.create(apellidos_nombres='PERSONA FICTICIA Pérez',dni='00123456')
        cls.periodo=PeriodoAsistencia.objects.get(titulo='Octubre 2026')
        cls.doc=TipoDocumento.objects.get(codigo_externo='01')
        cls.vac=TipoOperacion.objects.create(nombre='Categoría configurable',grupo_especial='VACACIONES',requiere_fecha_final=True)
        cls.sind=TipoOperacion.objects.create(nombre='Otra categoría configurable',grupo_especial='SINDICAL',requiere_fecha_final=True)
        cls.normal=TipoOperacion.objects.create(nombre='Vacaciones en nombre solamente',grupo_especial='NINGUNO',requiere_fecha_final=False)

    def registro(self,op=None,**extra):
        values=dict(trabajador=self.worker,periodo=self.periodo,tipo_documento=self.doc,tipo_operacion=op or self.vac,
            fecha_movimiento=date(2026,10,1),fecha_final=date(2026,10,3),hora_salida=time(9,15),
            cud='00-ABC',created_by=self.user,updated_by=self.user)
        values.update(extra);return Justificacion.objects.create(**values)

    def xlsx(self,response):
        self.assertEqual(response.status_code,200)
        return load_workbook(BytesIO(b''.join(response.streaming_content)),data_only=False)

    def test_duracion_inclusiva_sin_fecha_no_se_cuenta(self):
        self.assertEqual(self.registro().duracion,3)
        self.assertEqual(self.registro(fecha_final=date(2026,10,1)).duracion,1)
        self.assertIsNone(self.registro(fecha_final=None).duracion)
        obj=Justificacion(fecha_movimiento=date(2026,10,5),fecha_final=date(2026,10,1))
        self.assertIsNone(obj.duracion)

    def test_grupos_no_nombres_y_permiso_lectura_equivalente(self):
        self.registro(self.vac);self.registro(self.sind);self.registro(self.normal)
        self.client.force_login(self.reader)
        response=self.client.get(reverse('vacaciones'))
        self.assertEqual(response.context['pagina'].paginator.count,1)
        self.assertContains(response,'Categoría configurable');self.assertNotContains(response,'Vacaciones en nombre solamente')
        response=self.client.get(reverse('sindicales'))
        self.assertEqual(response.context['pagina'].paginator.count,1)
        self.assertContains(response,'Otra categoría configurable')
        self.assertNotContains(response,reverse('exportar_sindicales'))

    def test_filtros_sindical_fecha_inicial_final_y_periodo(self):
        esperado=self.registro(self.sind)
        self.registro(self.sind,fecha_movimiento=date(2026,9,1),fecha_final=date(2026,9,10))
        self.client.force_login(self.reader)
        response=self.client.get(reverse('sindicales'),{'q':'EREZ','dni':'0012','desde':'2026-10-01','hasta':'2026-10-09',
            'final_desde':'2026-10-03','final_hasta':'2026-10-03','periodo':self.periodo.pk})
        self.assertEqual(response.context['pagina'].paginator.count,1)
        self.assertEqual(response.context['pagina'][0].pk,esperado.pk)
        response=self.client.get(reverse('vacaciones'),{'final_desde':'2026-10-09','final_hasta':'2026-10-01'})
        self.assertTrue(response.context['filtros'].errors)
        self.assertEqual(response.context['pagina'].paginator.count,0)

    def test_exportaciones_respetan_filtros_y_todas_paginas(self):
        self.client.force_login(self.user)
        for n in range(30): self.registro(cud=f'00-COINCIDE-{n}')
        self.registro(self.sind,cud='EXCLUIDO')
        for ruta in ('exportar_consultas','exportar_vacaciones'):
            wb=self.xlsx(self.client.get(reverse(ruta),{'cud':'COINCIDE','page':2}))
            ws=wb.active;self.assertEqual(ws.max_row,31)
            self.assertEqual(ws.cell(2,3).value,'00123456')
            self.assertEqual(ws.cell(2,10).number_format,'dd/mm/yyyy')
            self.assertEqual(ws.cell(2,13).number_format,'hh:mm')
            self.assertEqual(ws.cell(2,12).value,3)
        wb=self.xlsx(self.client.get(reverse('exportar_sindicales')))
        self.assertEqual(wb.active.max_row,2)
        self.assertEqual(wb.active.cell(2,15).value,'EXCLUIDO')

    def test_exportacion_periodos_filtra_y_no_incluye_pii(self):
        self.client.force_login(self.user)
        wb=self.xlsx(self.client.get(reverse('exportar_periodos'),{'q':'Noviembre','estado':'ABIERTO'}))
        self.assertEqual(wb.active.max_row,2);self.assertEqual(wb.active.cell(2,1).value,'Noviembre 2026')
        self.assertIsNone(wb.active.cell(2,3).value)
        self.assertNotIn('DNI',[cell.value for cell in wb.active[1]])

    def test_exportacion_formulas_se_guardan_como_texto(self):
        self.client.force_login(self.user)
        dangerous=['=HYPERLINK("https://example.com","x")','+SUM(1,2)','-1+2','@SUM(1,2)','\t=1+2']
        for text in dangerous: self.registro(observaciones=text,cud=text)
        wb=self.xlsx(self.client.get(reverse('exportar_consultas')))
        rows=list(wb.active.iter_rows(min_row=2))
        self.assertEqual(len(rows),len(dangerous))
        self.assertEqual({row[15].value for row in rows},set(dangerous))
        for row in rows:
            self.assertEqual(row[14].data_type,'s');self.assertEqual(row[15].data_type,'s')
        self.assertFalse(any(cell.data_type=='f' for row in wb.active for cell in row))

    def test_permisos_exportacion_no_se_deducen_de_lectura(self):
        self.client.force_login(self.reader)
        for ruta in ('exportar_consultas','exportar_vacaciones','exportar_sindicales','exportar_periodos'):
            self.assertEqual(self.client.get(reverse(ruta)).status_code,403)
        self.reader.user_permissions.add(Permission.objects.get(codename='exportar_resultados'))
        self.client.force_login(self.reader)
        response=self.client.get(reverse('exportar_vacaciones'));self.assertEqual(response.status_code,200)
        # Consumir la descarga permite al cliente Django cerrar la respuesta
        # sin disparar close_old_connections dentro del atomic de TestCase.
        self.assertTrue(b''.join(response.streaming_content))
        u=User.objects.create_user('solo_exportar');u.user_permissions.add(Permission.objects.get(codename='exportar_resultados'))
        self.client.force_login(u)
        self.assertEqual(self.client.get(reverse('exportar_consultas')).status_code,403)

    def test_errores_de_filtros_exportan_solo_encabezados(self):
        self.registro();self.client.force_login(self.user)
        wb=self.xlsx(self.client.get(reverse('exportar_consultas'),{'fecha':'invalida'}))
        self.assertEqual(wb.active.max_row,1)

    def test_sin_autenticacion_no_hay_lectura_ni_descarga(self):
        for ruta in ('vacaciones','sindicales','exportar_consultas','exportar_periodos'):
            self.assertEqual(self.client.get(reverse(ruta)).status_code,302)
